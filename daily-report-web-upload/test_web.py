"""Verify Web-only privacy, session isolation, and in-memory downloads."""

import csv
import io
import os
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path
from unittest.mock import patch

from docx import Document
from streamlit.testing.v1 import AppTest

from dictionaries import load_dictionaries
from report_generator import FIELDS, WorkItem, generate_report
from web_config import get_deepseek_api_key


BASE = Path(__file__).resolve().parent


def start_comma_session(text):
    app = AppTest.from_file(str(BASE / "app.py"), default_timeout=30).run()
    app.radio[0].set_value("逗号文本（无需API）").run()
    app.text_area[0].set_value(text).run()
    next(button for button in app.button if button.label == "解析并预览").click().run()
    return app


class WebTests(unittest.TestCase):
    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": "environment-test-value"})
    @patch("web_config.st.secrets", {"DEEPSEEK_API_KEY": "secrets-test-value"})
    def test_environment_key_takes_priority(self):
        self.assertEqual(get_deepseek_api_key(), "environment-test-value")

    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""})
    @patch("web_config.st.secrets", {"DEEPSEEK_API_KEY": "secrets-test-value"})
    def test_streamlit_secrets_fallback(self):
        self.assertEqual(get_deepseek_api_key(), "secrets-test-value")

    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""})
    @patch("web_config.st.secrets", {})
    def test_missing_environment_and_secret(self):
        self.assertEqual(get_deepseek_api_key(), "")

    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""})
    def test_missing_key_keeps_comma_mode_available(self):
        app = AppTest.from_file(str(BASE / "app.py"), default_timeout=30).run()
        self.assertFalse(app.exception)
        self.assertTrue(any("服务器尚未配置" in message.value for message in app.info))
        app.text_area[0].set_value("热工今日工作：检查1#油泵").run()
        next(widget for widget in app.checkbox if widget.label.startswith("我同意")).check().run()
        app.button[0].click().run()
        self.assertTrue(any("服务器尚未配置" in message.value for message in app.error))
        self.assertNotIn("records", app.session_state)

        app = start_comma_session("热工,检查1#油泵,,,张三,,,")
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state["records"]), 1)
        self.assertEqual(app.session_state["records"][0]["人员"], "张三")

    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": "web-test-secret-never-expose-123456"})
    def test_server_key_is_not_rendered(self):
        app = AppTest.from_file(str(BASE / "app.py"), default_timeout=30).run()
        self.assertFalse(app.exception)
        visible = []
        for kind in ("title", "caption", "info", "warning", "error", "text", "markdown", "text_input"):
            visible.extend(str(widget.value) for widget in app.get(kind))
        self.assertNotIn(os.environ["DEEPSEEK_API_KEY"], "\n".join(visible))
        self.assertFalse(any(widget.label == "DeepSeek API Key" for widget in app.text_input))

    @patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""})
    def test_independent_sessions_do_not_share_reports(self):
        first = start_comma_session("热工,检查1#油泵,,,张三,,,")
        second = start_comma_session("机务,更换2#设备,,,李四,,,")
        self.assertFalse(first.exception)
        self.assertFalse(second.exception)
        self.assertEqual(first.session_state["records"][0]["人员"], "张三")
        self.assertEqual(second.session_state["records"][0]["人员"], "李四")
        self.assertNotEqual(first.session_state["records"], second.session_state["records"])

    def test_word_csv_bytes_and_relative_assets(self):
        item = WorkItem("机务", "检查240T1#", "否", "已完成", "张三", "更换75T，王五、赵六", "动火", "")
        with tempfile.TemporaryDirectory() as temporary:
            previous = Path.cwd()
            try:
                os.chdir(temporary)
                self.assertIn("professions", load_dictionaries())
                output = io.BytesIO()
                generate_report(BASE / "template" / "生产工作日报模板.docx", [item], date(2026, 9, 28), output)
            finally:
                os.chdir(previous)
        word_bytes = output.getvalue()
        self.assertTrue(word_bytes.startswith(b"PK"))
        table = Document(io.BytesIO(word_bytes)).tables[0]
        self.assertIn("更换75T，王五、赵六", table.rows[2].cells[8].text)
        self.assertEqual(table.rows[2].cells[9].text, "动火")

        csv_output = io.StringIO()
        writer = csv.writer(csv_output)
        writer.writerow(FIELDS)
        writer.writerow((item.profession, item.today, item.today_hazard, item.status,
                         item.people, item.tomorrow, item.tomorrow_hazard, item.remark))
        csv_bytes = csv_output.getvalue().encode("utf-8-sig")
        self.assertEqual(len(list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))), 2)
        self.assertNotIn(b"needs_review", csv_bytes)

    def test_concurrent_word_requests_use_separate_buffers(self):
        def create_word(name):
            item = WorkItem("机务", "检查" + name, "否", "已完成", name, "", "", "")
            output = io.BytesIO()
            generate_report(BASE / "template" / "生产工作日报模板.docx", [item], date(2026, 9, 28), output)
            return output.getvalue()

        with ThreadPoolExecutor(max_workers=2) as pool:
            first, second = list(pool.map(create_word, ("张三", "李四")))
        self.assertIn("检查张三", Document(io.BytesIO(first)).tables[0].rows[2].cells[3].text)
        self.assertIn("检查李四", Document(io.BytesIO(second)).tables[0].rows[2].cells[3].text)
        self.assertNotEqual(first, second)


if __name__ == "__main__":
    unittest.main()
