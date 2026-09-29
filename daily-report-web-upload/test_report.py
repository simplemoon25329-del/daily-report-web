import io
import json
import unittest
from pathlib import Path
from datetime import date
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn
from report_generator import WorkItem, align_work_items, parse_comma_text, generate_report, apply_defaults
from ai_service import validate_payload, extract_work

BASE = Path(__file__).resolve().parent

class Tests(unittest.TestCase):
    def item(self, **kw):
        d = dict(profession="机务", today="检查", today_hazard="", status="", people="张三",
                 tomorrow="", tomorrow_hazard="", remark="")
        d.update(kw)
        return WorkItem(**d)

    def render(self, rows):
        out = io.BytesIO()
        generate_report(BASE/"template/生产工作日报模板.docx", rows, date(2026,9,10), out)
        return Document(io.BytesIO(out.getvalue()))

    def test_csv(self):
        self.assertEqual(parse_comma_text('机务,"检查A,B",,,张三,,,')[0].today, "检查A,B")
        with self.assertRaises(ValueError):
            parse_comma_text("机务,缺字段")

    def test_layout(self):
        d = self.render([self.item(profession="热工"), self.item(remark="油管接近破损，需更换"),
                         self.item(tomorrow="任务一"), self.item(tomorrow="任务二")])
        t = d.tables[0]
        self.assertEqual(t.rows[2].cells[1].text.replace("\n",""), "机务")
        self.assertEqual(len(t.rows[2].cells[1].paragraphs), 1)
        self.assertIn("油管接近破损", t.rows[2].cells[3].text)
        self.assertEqual(t.rows[2].cells[10].text, "")
        self.assertEqual([t.rows[i].cells[7].text for i in (2,3)], ["1","2"])
        for i, size in [(2,"18"),(3,"4"),(4,"4"),(5,"18")]:
            border=t.rows[i].cells[3]._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:top"))
            self.assertEqual(border.get(qn("w:sz")), size)
        original=Document(BASE/"template/生产工作日报模板.docx")
        self.assertEqual([c.width for c in t.columns], [c.width for c in original.tables[0].columns])
        self.assertEqual(d.sections[0]._sectPr.xml, original.sections[0]._sectPr.xml)
        self.assertIn("09月10日",d.paragraphs[1].text)

    def test_today_tomorrow_alignment(self):
        today = [self.item(today=f"今日{i}", today_hazard=f"今危{i}", people=f"人员{i}")
                 for i in range(1, 5)]
        tomorrow = [self.item(today="", tomorrow=f"明日{i}", tomorrow_hazard=f"明危{i}", people="")
                    for i in range(1, 3)]
        aligned = align_work_items(today + tomorrow)
        self.assertEqual([(item.today, item.tomorrow) for item in aligned],
                         [("今日1", "明日1"), ("今日2", "明日2"), ("今日3", ""), ("今日4", "")])
        self.assertEqual([(item.today_hazard, item.people, item.tomorrow_hazard) for item in aligned],
                         [("今危1", "人员1", "明危1"), ("今危2", "人员2", "明危2"),
                          ("今危3", "人员3", ""), ("今危4", "人员4", "")])

        aligned = align_work_items(today[:2] + [
            self.item(today="", tomorrow=f"明日{i}", tomorrow_hazard=f"明危{i}", people="")
            for i in range(1, 5)
        ])
        self.assertEqual([(item.today, item.tomorrow) for item in aligned],
                         [("今日1", "明日1"), ("今日2", "明日2"), ("", "明日3"), ("", "明日4")])

    def test_expansion(self):
        d=self.render([self.item(today=f"任务{i}") for i in range(40)])
        self.assertEqual(len(d.tables[0].rows),42)
        self.assertEqual(d.tables[0].rows[-1].cells[2].text,"40")
        self.assertEqual(d.tables[0].rows[2].cells[1].text.replace("\n",""),"机务")

    def test_api_validation(self):
        payload={"professions":[{"profession":"机务","today_items":[
            {"today":"检查设备","people":"张三"}],"tomorrow_items":[]}],"warnings":[]}
        validated = validate_payload(payload)[0][0]
        self.assertEqual((validated.status, validated.today_hazard, validated.remark),
                         ("已完成", "否", ""))
        with self.assertRaises(ValueError):
            validate_payload({"professions":[{"profession":"机务","today_items":[],
                                                "tomorrow_items":"错误"}]})
        with self.assertRaises(ValueError):
            extract_work("test","")

    def test_api_structured_items_keep_equipment_numbers(self):
        payload = {"professions": [{
            "profession": "热工",
            "today_items": [
                {"today": "拆除240T1#炉压力变送器", "people": "徐超 赵千喜"},
                {"today": "检查2#首站设备", "people": "张三"},
                {"today": "更换75T滤芯", "people": "米中玉 沈一凡"},
                {"today": "校验25MW机组仪表", "people": "李四"},
            ],
            "tomorrow_items": [
                {"tomorrow": "恢复240T1#炉线路"},
                {"tomorrow": "复查2#首站"},
            ],
        }], "warnings": []}
        items, warnings = validate_payload(payload)
        aligned = align_work_items(items)
        self.assertEqual(len(items), 6)
        self.assertEqual(len(aligned), 4)
        self.assertEqual([(item.today, item.tomorrow) for item in aligned], [
            ("拆除240T1#炉压力变送器", "恢复240T1#炉线路"),
            ("检查2#首站设备", "复查2#首站"),
            ("更换75T滤芯", ""),
            ("校验25MW机组仪表", ""),
        ])
        self.assertEqual([item.people for item in aligned], ["徐超 赵千喜", "张三", "米中玉 沈一凡", "李四"])
        self.assertEqual([item.today_hazard for item in aligned], ["否"] * 4)
        self.assertEqual([item.tomorrow_hazard for item in aligned[:2]], ["否"] * 2)
        self.assertEqual(warnings, [])

    def test_api_mock(self):
        stage1 = {"professions": [{"profession": "机务",
            "today_items": [{"id": "T1", "raw_text": "检查"}], "tomorrow_items": []}],
            "warnings": ["核对"]}
        stage2 = {"today_items": [{"id": "T1", "today": "检查", "people": [],
            "today_hazard": None, "status": None, "remark": ""}], "tomorrow_items": []}
        responses = [io.BytesIO(json.dumps({"choices": [{"finish_reason": "stop",
            "message": {"content": json.dumps(value)}}]}).encode("utf-8"))
            for value in (stage1, stage2)]
        with patch("urllib.request.build_opener") as opener:
            opener.return_value.open.return_value.__enter__.side_effect = responses
            items,warnings=extract_work("检查", "test-only-key")
            self.assertEqual(warnings,["核对"])
            self.assertEqual(items[0].today,"检查")

    def test_defaults(self):
        default = apply_defaults(self.item(tomorrow="检查设备"))
        self.assertEqual((default.today_hazard, default.status, default.tomorrow_hazard, default.remark),
                         ("否", "已完成", "否", ""))
        explicit = self.item(today_hazard="动火",status="进行中",tomorrow="清理",
                             tomorrow_hazard="有限空间",remark="保留")
        self.assertEqual(apply_defaults(explicit), explicit)
        planned = apply_defaults(self.item(today="",tomorrow="清理"))
        self.assertEqual((planned.today_hazard,planned.status,planned.tomorrow_hazard),("","","否"))
        d=self.render([self.item()])
        c=d.tables[0].rows[2].cells
        self.assertEqual((c[4].text,c[5].text,c[9].text,c[10].text),("否","已完成","",""))

if __name__ == "__main__":
    unittest.main()
