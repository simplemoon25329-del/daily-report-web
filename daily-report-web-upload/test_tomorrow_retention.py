"""Regress tomorrow field retention through extraction, preview and export."""
from datetime import date
import io
import json
import os
from pathlib import Path
import unittest
from unittest.mock import patch

from docx import Document

from ai_service import extract_work
from report_generator import align_work_items, generate_report


class TomorrowRetentionTests(unittest.TestCase):
    def setUp(self):
        self.raw_plans = [
            '1、检查1#油泵，张三、李四',
            '2、更换75T设备，王五、赵六，动火作业',
            '3、进入罐体检查，程保华、左明显，有限空间作业',
            '4、处理2#设备问题，负责人张三，明日上午完成（厂房内，待验收，备注：携带工具）',
        ]
        self.plans = [
            '检查1#油泵，张三、李四',
            '更换75T设备，王五、赵六',
            '进入罐体检查，程保华、左明显',
            '处理2#设备问题，负责人张三，明日上午完成（厂房内，待验收，备注：携带工具）',
        ]
        self.today_raw = '检查240T1#，徐超、赵千喜，危险作业：动火作业，未完成，备注：等待验收'
        self.text = '热工今日工作：' + self.today_raw + '\n明日计划：\n' + '\n'.join(self.raw_plans)
        self.stage1 = {'professions': [{'profession': '热工',
            'today_items': [{'id': 'T1', 'raw_text': self.today_raw}],
            'tomorrow_items': [{'id': 'M' + str(index + 1), 'raw_text': raw}
                               for index, raw in enumerate(self.raw_plans)]}], 'warnings': []}
        self.stage2 = {'today_items': [{'id': 'T1', 'today': '检查240T1#',
            'people': ['徐超', '赵千喜'], 'today_hazard': '动火作业',
            'status': '未完成', 'remark': '等待验收'}],
            'tomorrow_items': [{'id': 'M' + str(index + 1), 'tomorrow': text,
                                'tomorrow_hazard': hazard}
                               for index, (text, hazard) in enumerate(zip(self.plans, [None, '动火', '有限', None]))],
            'warnings': []}

    def extract(self):
        details = []
        with patch('ai_service._request_completion', side_effect=[
            json.dumps(self.stage1, ensure_ascii=False), json.dumps(self.stage2, ensure_ascii=False)
        ]) as request:
            items, warnings = extract_work(self.text, 'test-only-key', review_details=details)
        self.assertEqual(request.call_count, 2)
        self.assertEqual(json.loads(request.call_args_list[1].args[1]), self.stage1)
        self.assertEqual(warnings, [])
        return align_work_items(items), details

    def test_names_without_hazard(self):
        rows, details = self.extract()
        self.assertIsNone(self.stage2['tomorrow_items'][0]['tomorrow_hazard'])
        self.assertEqual(rows[0].tomorrow, self.plans[0])
        self.assertEqual(rows[0].tomorrow_hazard, '否')
        self.assertEqual(details[1]['raw_text'], self.raw_plans[0])
        self.assertEqual(details[1]['result']['tomorrow'], self.plans[0])

    def test_hot_work_preserves_names(self):
        rows, _ = self.extract()
        self.assertEqual(rows[1].tomorrow, self.plans[1])
        self.assertNotIn('动火作业', rows[1].tomorrow)
        self.assertEqual(rows[1].tomorrow_hazard, '动火')

    def test_confined_space_preserves_names(self):
        rows, _ = self.extract()
        self.assertEqual(rows[2].tomorrow, self.plans[2])
        self.assertEqual(rows[2].tomorrow_hazard, '有限')

    def test_owner_time_status_location_and_parentheses(self):
        rows, _ = self.extract()
        self.assertEqual(rows[3].tomorrow, self.plans[3])
        self.assertEqual(rows[3].tomorrow_hazard, '否')

    def test_today_extraction_unchanged(self):
        rows, _ = self.extract()
        self.assertEqual((rows[0].today, rows[0].people, rows[0].today_hazard,
                          rows[0].status, rows[0].remark),
                         ('检查240T1#', '徐超 赵千喜', '动火作业', '未完成', '等待验收'))
        self.assertTrue(all(not row.people and not row.today for row in rows[1:]))
        self.assertEqual(len(rows), 4)

    @patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-only-key'})
    def test_preview_comparison_and_word_keep_tomorrow_text(self):
        from streamlit.testing.v1 import AppTest
        base = Path(__file__).parent
        app = AppTest.from_file(str(base / 'app.py'), default_timeout=30).run()
        app.text_area[0].set_value(self.text).run()
        next(widget for widget in app.checkbox if widget.label.startswith('我同意')).check().run()
        with patch('ai_service._request_completion', side_effect=[json.dumps(self.stage1), json.dumps(self.stage2)]):
            app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertEqual([row['明日计划'] for row in app.session_state['records']], self.plans)
        comparison = '\n'.join(widget.value for widget in app.text)
        for raw, plan in zip(self.raw_plans, self.plans):
            self.assertIn(raw, comparison)
            self.assertIn('明日计划：' + plan, comparison)
        output = io.BytesIO()
        generate_report(base / 'template/生产工作日报模板.docx', self.extract()[0], date(2026, 9, 22), output)
        table = Document(output).tables[0]
        self.assertEqual([row.cells[8].text for row in table.rows[2:]], self.plans)
        self.assertEqual([row.cells[9].text for row in table.rows[2:]], ['否', '动火', '有限', '否'])
        for index in (2, 5):
            for paragraph in table.rows[index].cells[9].paragraphs:
                for run in paragraph.runs:
                    self.assertFalse(run.bold)
                    self.assertEqual(str(run.font.color.rgb), '000000')


if __name__ == '__main__':
    unittest.main()
